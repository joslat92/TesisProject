"""Apply the audited, deterministic corrections to the final thesis DOCX.

This script intentionally targets verified paragraph prefixes and table cells.
It fails when the expected baseline structure is not present instead of making
an ambiguous edit.
"""

from __future__ import annotations

import argparse
import os
import shutil
import tempfile
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DOCX = ROOT / "Tesis Maestro Final.docx"
MZ_FIGURE = ROOT / "reports" / "figs" / "Fig_T20_calibracion_scatter.png"


def normalized(text: str) -> str:
    return " ".join(text.split())


def matching_paragraphs(doc: Document, prefix: str) -> list[Paragraph]:
    return [p for p in doc.paragraphs if normalized(p.text).startswith(prefix)]


def find_one(doc: Document, prefix: str) -> Paragraph:
    matches = matching_paragraphs(doc, prefix)
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph starting with {prefix!r}; found {len(matches)}"
        )
    return matches[0]


def find_one_with_style(doc: Document, prefix: str, style_name: str) -> Paragraph:
    matches = [
        p for p in matching_paragraphs(doc, prefix)
        if p.style is not None and p.style.name == style_name
    ]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one {style_name!r} paragraph starting with {prefix!r}; "
            f"found {len(matches)}"
        )
    return matches[0]


def set_paragraph(doc: Document, prefix: str, text: str) -> None:
    find_one(doc, prefix).text = text


def set_paragraph_once(doc: Document, old_prefix: str, text: str) -> None:
    matches = matching_paragraphs(doc, old_prefix)
    if len(matches) == 1:
        matches[0].text = text
        return
    if not matches and any(normalized(p.text) == normalized(text) for p in doc.paragraphs):
        return
    raise RuntimeError(
        f"Could not apply idempotent paragraph edit for {old_prefix!r}; "
        f"found {len(matches)} baseline matches"
    )


def insert_paragraph_before(paragraph: Paragraph, text: str) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addprevious(element)
    inserted = Paragraph(element, paragraph._parent)
    inserted.style = paragraph.style
    inserted.add_run(text)
    return inserted


def replace_in_all_paragraphs(doc: Document, old: str, new: str) -> int:
    count = 0
    for paragraph in doc.paragraphs:
        hits = paragraph.text.count(old)
        if hits:
            paragraph.text = paragraph.text.replace(old, new)
            count += hits
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    hits = paragraph.text.count(old)
                    if hits:
                        paragraph.text = paragraph.text.replace(old, new)
                        count += hits
    return count


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.text = text
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def update_result_tables(doc: Document) -> None:
    dm = doc.tables[6]
    expected = {
        (3, 4): "1.87 (p=0.062)",
        (4, 2): "1.98 (p=0.049)",
        (5, 4): "1.80 (p=0.073)",
    }
    replacements = {
        (3, 4): "1.89 (p=0.060)",
        (4, 2): "1.99 (p=0.048)",
        (5, 4): "1.82 (p=0.070)",
    }
    for key, old in expected.items():
        row, col = key
        actual = normalized(dm.cell(row, col).text)
        if actual != old:
            raise RuntimeError(f"Unexpected Table 7.2 value at {key}: {actual!r}")
        set_cell_text(dm.cell(row, col), replacements[key])

    glossary = doc.tables[12]
    if normalized(glossary.cell(1, 0).text) != "Walk-forward validation":
        raise RuntimeError("Glossary table is not in the expected position")
    set_cell_text(
        glossary.cell(1, 1),
        "Validación temporal por bloques mensuales: cada modelo se re-estima "
        "con una ventana expanding que termina antes del bloque evaluado.",
    )


def append_reference(table, number: str, reference: str) -> None:
    new_row = table.add_row()
    if len(new_row.cells) != 2:
        raise RuntimeError("Reference table no longer has two columns")
    source = table.rows[-2]
    for source_cell, target_cell in zip(source.cells, new_row.cells):
        existing_tc_pr = target_cell._tc.tcPr
        if existing_tc_pr is not None:
            target_cell._tc.remove(existing_tc_pr)
        source_tc_pr = source_cell._tc.tcPr
        if source_tc_pr is not None:
            target_cell._tc.insert(0, deepcopy(source_tc_pr))
    set_cell_text(new_row.cells[0], number)
    set_cell_text(new_row.cells[1], reference)


def update_references(doc: Document) -> None:
    references = doc.tables[10]
    if normalized(references.cell(0, 0).text) != "[1]":
        raise RuntimeError("Numbered bibliography table was not found")
    if len(references.rows) != 26:
        raise RuntimeError(f"Expected 26 bibliography entries; found {len(references.rows)}")

    append_reference(
        references,
        "[27]",
        "D. I. Harvey, S. J. Leybourne, y P. Newbold, «Testing the equality "
        "of prediction mean squared errors», International Journal of "
        "Forecasting, vol. 13, n.o 2, pp. 281–291, 1997, "
        "doi: 10.1016/S0169-2070(96)00719-4.",
    )
    append_reference(
        references,
        "[28]",
        "M. H. Pesaran y A. Timmermann, «A Simple Nonparametric Test of "
        "Predictive Performance», Journal of Business & Economic Statistics, "
        "vol. 10, n.o 4, pp. 461–465, 1992, "
        "doi: 10.1080/07350015.1992.10509922.",
    )
    append_reference(
        references,
        "[29]",
        "J. A. Mincer y V. Zarnowitz, «The Evaluation of Economic Forecasts», "
        "en Economic Forecasts and Expectations: Analysis of Forecasting "
        "Behavior and Performance, National Bureau of Economic Research, "
        "1969, pp. 3–46.",
    )


def update_methodology(doc: Document) -> None:
    set_paragraph(
        doc,
        "Se comparan ocho especificaciones:",
        "Se comparan siete especificaciones activas: Random Walk (RW), ARIMA, "
        "ARIMAX, SARIMAX, LSTM (plain), LSTM+Sent y LSTM+Sent+VIX. El gate "
        "estacional no activó términos m = 5, por lo que los artefactos "
        "SARIMAX coinciden con ARIMAX; no se reporta una especificación SARIMA "
        "independiente.",
    )
    set_paragraph(
        doc,
        "OOS 2024 completo:",
        "OOS 2024 completo: periodo de evaluación enero–diciembre de 2024. "
        "ARIMA, ARIMAX y SARIMAX usan ventana expanding con re-estimación de "
        "parámetros al inicio de cada mes y re-filtrado diario; las LSTM se "
        "entrenan con información hasta el 31-12-2023 y mantienen parámetros "
        "fijos durante el OOS.",
    )
    set_paragraph(
        doc,
        "Walk-forward mensual 2024:",
        "Walk-forward mensual 2024: doce bloques consecutivos; para cada mes "
        "m de 2024 se re-estima cada modelo con información hasta el cierre "
        "del mes m−1 y se predicen los días del mes m.",
    )
    set_paragraph(
        doc,
        "Estacionariedad. Sobre",
        "Estacionariedad. Las pruebas ADF y KPSS se aplican a Target_Price en "
        "nivel y primera diferencia, como se documenta en el Apéndice A. Los "
        "modelos clásicos se ajustan sobre el retorno logarítmico diario "
        "ret_1d, que ya es una primera diferencia de log(P); por ello usan "
        "d = 0 sobre retornos, equivalente a d = 1 sobre log(P).",
    )
    set_paragraph(
        doc,
        "se determina por criterio de información",
        "El diagnóstico justifica la transformación a retornos; no selecciona "
        "dinámicamente un valor de d distinto para cada horizonte.",
    )
    set_paragraph(
        doc,
        "Estacionalidad. Se aplican tres herramientas",
        "Estacionalidad. El pipeline aplica un único gate sobre los retornos "
        "diarios del periodo in-sample.",
    )
    set_paragraph(
        doc,
        "Espectro de potencia:",
        "Se agrupan los retornos por día de la semana y se aplica la prueba "
        "no paramétrica de Kruskal–Wallis.",
    )
    set_paragraph(
        doc,
        "Autocorrelación estacional:",
        "La hipótesis nula es igualdad de las distribuciones entre lunes y "
        "viernes; el gate sólo usa datos hasta el cierre de 2023.",
    )
    set_paragraph(
        doc,
        "Prueba tipo Fourier",
        "Regla de decisión: si p < 0.10 se habilita la especificación "
        "estacional m = 5; en caso contrario se usa m = 0.",
    )
    set_paragraph(
        doc,
        "Regla de decisión: si hay evidencia",
        "En los datos del estudio se obtuvo p = 0.302, por lo que los términos "
        "estacionales permanecieron desactivados en todos los horizontes.",
    )
    set_paragraph(
        doc,
        "ARIMA(p,d,q): identificación",
        "ARIMA: orden fijo (1,0,1) sobre ret_1d, equivalente a ARIMA(1,1,1) "
        "sobre log(P). El orden está declarado en config.yaml y no se "
        "selecciona mediante una rejilla AIC durante la ejecución.",
    )
    set_paragraph(
        doc,
        "SARIMA(p,d,q)",
        "Componente estacional: si el gate se activara, SARIMAX usaría el "
        "orden estacional configurado (0,1,0,5). No existe una salida SARIMA "
        "independiente entre los siete modelos activos.",
    )
    set_paragraph(
        doc,
        "SARIMAX: combinación",
        "SARIMAX: combina el orden no estacional (1,0,1) sobre ret_1d con las "
        "exógenas rezagadas y el componente estacional condicionado al gate. "
        "Como el gate no se activó, sus predicciones coinciden con ARIMAX.",
    )
    set_paragraph(
        doc,
        "En ARIMA/SARIMA (y variantes con X)",
        "Para cada fecha de emisión t, los modelos clásicos pronostican h "
        "retornos diarios de forma iterada y los suman. En OOS sus parámetros "
        "se re-estiman al inicio de cada mes; en WF se estiman una vez con "
        "datos hasta el cierre del bloque anterior y se re-filtran diariamente "
        "dentro del bloque.",
    )
    set_paragraph(
        doc,
        "Entrenamiento IS:",
        "Entrenamiento: en el OOS 2024 las LSTM se ajustan con datos hasta el "
        "31-12-2023, mientras los modelos clásicos usan re-estimación mensual "
        "expanding. En WF todos los modelos se re-estiman con datos hasta el "
        "bloque previo. En la robustez 2025 las LSTM se ajustan hasta el "
        "31-12-2024 y los modelos clásicos continúan con re-estimación mensual.",
    )
    set_paragraph(
        doc,
        "Predicción OOS/WF:",
        "Predicción OOS/WF: en cada origen t se generan h pasos diarios y se "
        "suman para obtener el retorno acumulado del horizonte.",
    )
    set_paragraph(
        doc,
        "OOS: preds/OOS/",
        "OOS: outputs/preds/OOS/preds_T{h}_{modelo}.csv, con Date, h, model, "
        "y_true_ret, y_pred_ret, y_true_level y y_pred_level.",
    )
    set_paragraph(
        doc,
        "WF: preds/WF/",
        "WF: outputs/preds/WF/preds_T{h}_{modelo}_block{m}.csv, con las mismas "
        "columnas y el identificador block.",
    )
    set_paragraph(
        doc,
        "Artefactos opcionales:",
        "Los parámetros, semillas e hiperparámetros quedan declarados en los "
        "archivos de configuración; el pipeline no persiste pesos neuronales "
        "como artefactos canónicos.",
    )
    set_paragraph(
        doc,
        "Se contrasta el modelo A contra",
        "Se contrasta el modelo A contra un comparador B usando diferencias "
        "de pérdidas cuadráticas con corrección HAC y regla L = h−1 [9]. Se "
        "reportan el estadístico original y la corrección de muestra pequeña "
        "de Harvey–Leybourne–Newbold [27].",
    )

    mz_heading = find_one(doc, "Prueba de Mincer–Zarnowitz (MZ) en niveles.")
    mz_explanation = mz_heading._p.getnext()
    while mz_explanation is not None and mz_explanation.tag != qn("w:p"):
        mz_explanation = mz_explanation.getnext()
    if mz_explanation is None:
        raise RuntimeError("MZ explanatory paragraph was not found")
    mz_paragraph = Paragraph(mz_explanation, mz_heading._parent)
    mz_paragraph.text = (
        "Se estima por OLS con covarianza HAC y máximo de rezagos h−1. La "
        "calibración se evalúa con dos contrastes separados, α = 0 y β = 1; "
        "se reportan α, β, R² y ambos p-valores [29]."
    )
    insert_paragraph_before(mz_paragraph, "MZ_EQUATION_PLACEHOLDER")

    set_paragraph(
        doc,
        "Scripts: src/stages/00_prepare.py",
        "Scripts: src/stages/00_prepare.py (features y splits); "
        "src/stages/10_baselines.py–16_robustez_2025.py (entrenamiento y "
        "validación); src/stages/20_evaluate_stats.py y 24_wf_metrics.py "
        "(métricas, DM y MZ); src/stages/30_make_figures.py y 31_regimes.py "
        "(figuras y análisis por régimen).",
    )
    set_paragraph(
        doc,
        "Pruebas: dm_summary.csv",
        "Pruebas: reports/data/tbl_DM_OOS.csv, tbl_MZ_core.csv, tbl_PT_OOS.csv "
        "y tbl_DM_2025.csv.",
    )
    set_paragraph(
        doc,
        "Selección estacional por horizonte",
        "Selección estacional: gate Kruskal–Wallis sobre el periodo in-sample; "
        "p < 0.10 habilita m = 5 y, en caso contrario, m = 0.",
    )
    set_paragraph(
        doc,
        "Alternativa de orden ARIMA:",
        "Orden ARIMA: (1,0,1) fijo sobre ret_1d, declarado en config.yaml y "
        "aplicado de forma uniforme en OOS, WF y robustez 2025.",
    )
    set_paragraph(
        doc,
        "Análisis por régimen (opcional):",
        "Análisis por régimen: segmentación del OOS por cuartiles del VIX "
        "contemporáneo en la fecha de origen t; los umbrales se calculan por "
        "horizonte y se aplican por igual a todos los modelos.",
    )


def update_results_and_appendices(doc: Document) -> None:
    set_paragraph(
        doc,
        "Las regresiones de Mincer–Zarnowitz en niveles",
        "Las regresiones de Mincer–Zarnowitz en niveles muestran un contraste "
        "nítido entre horizontes. En T+1, los siete modelos están bien "
        "calibrados: pendientes entre 0.983 y 0.989, interceptos no "
        "distinguibles de cero (p ≥ 0.08) y R² de 0.975. Desde T+5, las "
        "hipótesis α = 0 y β = 1 se rechazan por separado para todos los "
        "modelos, incluido el paseo aleatorio, con pendientes que descienden "
        "hasta 0.80–0.82 y R² de 0.70–0.71 en T+20. La uniformidad de ambos "
        "rechazos indica que la descalibración caracteriza al horizonte en el "
        "periodo evaluado y no discrimina entre familias de modelos.",
    )
    find_one_with_style(
        doc, "Figura 7.3 — Calibración Mincer–Zarnowitz", "Caption"
    ).text = (
        "Figura 7.3 — Calibración Mincer–Zarnowitz en T+20 (OOS 2024): nivel "
        "observado contra nivel pronosticado por modelo. La línea de 45° "
        "representa la calibración perfecta (α = 0, β = 1); los coeficientes "
        "anotados corresponden a la regresión en niveles."
    )
    set_paragraph(
        doc,
        "Como prueba de estrés de las conclusiones",
        "Como prueba de estrés, los modelos se evaluaron sobre enero–abril de "
        "2025, periodo que contiene una caída severa entre febrero y abril. "
        "Las LSTM se entrenaron con información hasta diciembre de 2024; los "
        "modelos clásicos usaron ventana expanding con re-estimación mensual "
        "durante el bloque. El RMSE de T+1 prácticamente se duplicó respecto "
        "de 2024 (de ~0.011 a ~0.022), y el detalle mensual muestra el salto "
        "del shock.",
    )
    find_one_with_style(
        doc, "Tabla 7.5 — Bloque de robustez:", "Caption"
    ).text = (
        "Tabla 7.5 — Bloque de robustez: enero–abril de 2025 (LSTM entrenadas "
        "hasta dic-2024; modelos clásicos con re-estimación mensual)"
    )
    set_paragraph(
        doc,
        "En series de índices bursátiles diarios",
        "Los resultados muestran no estacionariedad de Target_Price en nivel "
        "y estacionariedad después de una diferencia. En el pipeline, ARIMA y "
        "SARIMAX se ajustan con orden (1,0,1) sobre ret_1d, equivalente a "
        "ARIMA(1,1,1) sobre log(P). Las LSTM pronostican los horizontes "
        "h ∈ {1, 5, 10, 20} a partir de ventanas históricas de 40 días.",
    )
    set_paragraph(
        doc,
        "git clone --branch reestructura-dic2025",
        "git clone --branch correcciones-auditoria-final "
        "https://github.com/joslat92/TesisProject.git\n"
        "cd TesisProject\n"
        "python -m venv .venv\n"
        ".venv\\Scripts\\python.exe -m pip install --upgrade pip\n"
        ".venv\\Scripts\\python.exe -m pip install -r requirements.txt",
    )
    set_paragraph(
        doc,
        "El código se organiza en dos paquetes bajo src/.",
        "El código se organiza bajo src/. El núcleo (src/core/) contiene el "
        "validador de contrato con fail-fast (contract.py), la prueba de "
        "Diebold–Mariano con kernel de Bartlett (dm.py) y las métricas "
        "compartidas (metrics.py). Las etapas ejecutables se ubican en "
        "src/stages/ y se orquestan mediante run_all.py.",
    )
    set_paragraph(
        doc,
        "00_prepare.py Preparación de datos",
        "00_prepare.py  Preparación de datos (parquets por horizonte)\n"
        "10_baselines.py  Baseline RW\n"
        "11_train_arima.py  ARIMA iterado\n"
        "13_train_sarimax.py  ARIMAX/SARIMAX con exógenas congeladas\n"
        "12_train_lstm.py  Variantes LSTM (OOS 2024)\n"
        "14_walkforward.py  Walk-forward: 12 bloques × 7 modelos\n"
        "15_multiseed_lstm.py  Multi-semilla LSTM (120 entrenamientos)\n"
        "16_robustez_2025.py  Bloque de robustez enero–abril 2025\n"
        "24_wf_metrics.py  Métricas walk-forward\n"
        "20_evaluate_stats.py  Evaluación consolidada y gate completo\n"
        "30_make_figures.py  Figuras canónicas\n"
        "31_regimes.py  Regímenes por cuartiles de VIX",
    )
    set_paragraph(
        doc,
        "# Verificación rápida de mecánica",
        "# Verificación rápida de mecánica en copia temporal aislada\n"
        ".venv\\Scripts\\python.exe run_all.py --quick\n\n"
        "# Reproducción completa (varias horas: incluye walk-forward\n"
        "# y multi-semilla)\n"
        ".venv\\Scripts\\python.exe run_all.py",
    )
    set_paragraph(
        doc,
        "En el modo --quick los modelos clásicos",
        "El modo --quick crea una copia temporal aislada del repositorio, "
        "reduce las LSTM a dos épocas y omite walk-forward, sus métricas y el "
        "análisis multi-semilla. Sus resultados neuronales sólo verifican la "
        "mecánica y no corresponden a las cifras de la tesis; los artefactos "
        "canónicos del árbol de trabajo permanecen intactos. La reproducción "
        "completa, sin banderas, regenera todos los resultados.",
    )


def update_final_polish(doc: Document) -> None:
    set_paragraph_once(
        doc,
        "SARIMA/SARIMAX pueden no converger",
        "El componente estacional sólo se estima si el gate lo habilita; cuando "
        "m = 0, SARIMAX coincide por construcción con ARIMAX.",
    )
    set_paragraph_once(
        doc,
        "Diagnóstico ADF/KPSS y batería de estacionalidad",
        "Diagnóstico ADF/KPSS sobre nivel y primera diferencia; gate estacional "
        "Kruskal–Wallis sobre el periodo in-sample.",
    )
    set_paragraph_once(
        doc,
        "Elección de m (5 o 0)",
        "Elección de m ∈ {5, 0} mediante el gate; orden ARIMA (1,0,1) fijo "
        "sobre ret_1d según config.yaml.",
    )
    set_paragraph_once(
        doc,
        "Entrenamiento de RW / ARIMA / ARIMAX / SARIMA",
        "Entrenamiento de los siete modelos activos: RW, ARIMA, ARIMAX, "
        "SARIMAX, LSTM, LSTM+Sent y LSTM+Sent+VIX.",
    )
    set_paragraph_once(
        doc,
        "Generación de predicciones OOS 2024",
        "Generación de predicciones OOS 2024, WF mensual y robustez 2025; "
        "validación del contrato y persistencia de artefactos.",
    )
    set_paragraph_once(
        doc,
        "Cálculo de MAE, RMSE, MDA, DM",
        "Cálculo de MAE, RMSE, MDA, DM (HAC L = h−1 y corrección HLN), MZ "
        "(HAC) y PT; generación de figuras y tablas.",
    )

    appendix_heading = find_one(doc, "APÉNDICE")
    appendix_heading.paragraph_format.page_break_before = True
    appendix_heading.paragraph_format.keep_with_next = True
    find_one_with_style(
        doc, "Apéndice A. Pruebas de estacionariedad", "Caption"
    ).paragraph_format.keep_with_next = True
    find_one_with_style(
        doc, "Apéndice B. Reproducibilidad", "Heading 1"
    ).paragraph_format.keep_with_next = True
    find_one_with_style(
        doc, "Apéndice C. Glosario técnico", "Caption"
    ).paragraph_format.keep_with_next = True

    final_table = doc.tables[12]._tbl
    element = final_table.getnext()
    while element is not None and element.tag != qn("w:sectPr"):
        next_element = element.getnext()
        if element.tag != qn("w:p") or normalized(
            "".join(element.xpath('.//*[local-name()="t"]/text()'))
        ):
            raise RuntimeError("Unexpected content after the final glossary table")
        element.getparent().remove(element)
        element = next_element


def prepare_equation_placeholders(doc: Document) -> None:
    adf_caption = find_one(doc, "Ecuación 1. Ecuación ADF")
    adf_equation = adf_caption._p.getprevious()
    if adf_equation is None or adf_equation.tag != qn("w:p"):
        raise RuntimeError("ADF equation paragraph was not found")
    Paragraph(adf_equation, adf_caption._parent).text = "ADF_EQUATION_PLACEHOLDER"

    set_paragraph(
        doc,
        "Se calcula el estadístico",
        "Se calcula el estadístico τ_ADF = γ̂/SE(γ̂) y se contrasta con valores "
        "críticos no estándar (Dickey y Fuller, 1979 [7]). Rechazar H₀ implica "
        "que la serie es estacionaria en nivel o en la diferencia utilizada.",
    )

    kpss_caption = find_one(doc, "Ecuación 2. Estadístico KPSS")
    kpss_equation = kpss_caption._p.getprevious()
    if kpss_equation is None or kpss_equation.tag != qn("w:p"):
        raise RuntimeError("KPSS equation paragraph was not found")
    Paragraph(kpss_equation, kpss_caption._parent).text = "KPSS_EQUATION_PLACEHOLDER"


def refresh_all_equation_placeholders(doc: Document) -> None:
    prepare_equation_placeholders(doc)
    mz_heading = find_one(doc, "Prueba de Mincer–Zarnowitz (MZ) en niveles.")
    mz_equation = mz_heading._p.getnext()
    while mz_equation is not None and mz_equation.tag != qn("w:p"):
        mz_equation = mz_equation.getnext()
    if mz_equation is None:
        raise RuntimeError("MZ equation paragraph was not found")
    Paragraph(mz_equation, mz_heading._parent).text = "MZ_EQUATION_PLACEHOLDER"


def update_numeric_values(doc: Document) -> None:
    replacements = [
        ("p = 0.062–0.073", "p = 0.060–0.070"),
        ("p = 0.062 y p = 0.073", "p = 0.060 y p = 0.070"),
        ("p = 0.049", "p = 0.048"),
        ("p = 0.062", "p = 0.060"),
        ("p = 0.073", "p = 0.070"),
        ("mediana de p-valores 0.194", "mediana de p-valores 0.190"),
        ("mediana de los p-valores fue 0.194", "mediana de los p-valores fue 0.190"),
        ("median p-value 0.194", "median p-value 0.190"),
        ("ensemble p = 0.140", "ensemble p = 0.136"),
        ("p = 0.140", "p = 0.136"),
        ("p = 0.042", "p = 0.041"),
        ("DM = −2.10", "DM = −2.12"),
        ("p = 0.039", "p = 0.038"),
        ("p = 0.015–0.029", "p = 0.009–0.018"),
        ("[Makridakis et al.]", "[5]"),
    ]
    for old, new in replacements:
        replace_in_all_paragraphs(doc, old, new)


def add_method_citations(doc: Document) -> None:
    pt = find_one(doc, "La prueba de Pesaran–Timmermann sobre")
    if "[28]" not in pt.text:
        pt.text = pt.text.replace(
            "La prueba de Pesaran–Timmermann",
            "La prueba de Pesaran–Timmermann [28]",
            1,
        )


def figure_relationship_target(doc: Document) -> str:
    caption = find_one_with_style(
        doc, "Figura 7.3 — Calibración Mincer–Zarnowitz", "Caption"
    )
    element = caption._p.getprevious()
    while element is not None:
        blips = element.xpath(".//a:blip")
        if blips:
            rid = blips[0].get(qn("r:embed"))
            return doc.part.rels[rid].target_ref
        element = element.getprevious()
    raise RuntimeError("Embedded image preceding Figure 7.3 was not found")


def replace_zip_member(docx: Path, member: str, source: Path) -> None:
    if not source.exists():
        raise FileNotFoundError(source)
    fd, temp_name = tempfile.mkstemp(suffix=".docx", dir=docx.parent)
    os.close(fd)
    temp = Path(temp_name)
    try:
        found = False
        with zipfile.ZipFile(docx, "r") as src, zipfile.ZipFile(
            temp, "w", compression=zipfile.ZIP_DEFLATED
        ) as dst:
            for item in src.infolist():
                data = src.read(item.filename)
                if item.filename == member:
                    data = source.read_bytes()
                    found = True
                dst.writestr(item, data)
        if not found:
            raise RuntimeError(f"DOCX member {member!r} was not found")
        shutil.move(temp, docx)
    finally:
        temp.unlink(missing_ok=True)


def save_document(doc: Document, source: Path, output: Path) -> None:
    image_target = figure_relationship_target(doc)
    member = "word/" + image_target.replace("\\", "/")
    output.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(suffix=".docx", dir=output.parent)
    os.close(fd)
    temp = Path(temp_name)
    try:
        doc.save(temp)
        replace_zip_member(temp, member, MZ_FIGURE)
        shutil.move(temp, output)
    finally:
        temp.unlink(missing_ok=True)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=DEFAULT_DOCX)
    parser.add_argument("--output", type=Path, default=DEFAULT_DOCX)
    parser.add_argument(
        "--post-word",
        action="store_true",
        help="apply idempotent textual fixes after Word equation conversion",
    )
    parser.add_argument(
        "--refresh-equations",
        action="store_true",
        help="replace the three audited equations with Word placeholders",
    )
    args = parser.parse_args()

    source = args.input.resolve()
    output = args.output.resolve()
    doc = Document(source)
    if len(doc.tables) != 13:
        raise RuntimeError(f"Expected 13 tables in baseline DOCX; found {len(doc.tables)}")

    if args.post_word:
        update_numeric_values(doc)
        update_final_polish(doc)
        save_document(doc, source, output)
        print(f"Post-Word fixes written to: {output}")
        return

    if args.refresh_equations:
        refresh_all_equation_placeholders(doc)
        save_document(doc, source, output)
        print(f"Equation placeholders written to: {output}")
        return

    update_result_tables(doc)
    update_methodology(doc)
    update_results_and_appendices(doc)
    update_final_polish(doc)
    prepare_equation_placeholders(doc)
    update_numeric_values(doc)
    add_method_citations(doc)
    update_references(doc)
    save_document(doc, source, output)
    print(f"Updated thesis written to: {output}")


if __name__ == "__main__":
    main()
