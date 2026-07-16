"""Update the thesis with the sealed results from the reconstructed dataset."""

from __future__ import annotations

import argparse
from copy import deepcopy
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "Tesis Maestro Final.docx"
DEFAULT_OUTPUT = ROOT / "Tesis Maestro Final - Datos Reconstruidos.docx"


MODEL_LABELS = {
    "RW": "RW (paseo aleatorio)",
    "ARIMA": "ARIMA",
    "ARIMAX": "ARIMAX (≡ SARIMAX)",
    "LSTM": "LSTM",
    "LSTM_SENT": "LSTM+Sent",
    "LSTM_FULL": "LSTM+Sent+VIX",
}
MODEL_ORDER = ["RW", "ARIMA", "ARIMAX", "LSTM", "LSTM_SENT", "LSTM_FULL"]
LSTM_ORDER = ["LSTM", "LSTM_SENT", "LSTM_FULL"]


def normalized(value: str) -> str:
    return " ".join(value.split())


def find_one(doc: Document, prefix: str, style: str | None = None) -> Paragraph:
    matches = [
        paragraph
        for paragraph in doc.paragraphs
        if normalized(paragraph.text).startswith(normalized(prefix))
        and (style is None or paragraph.style.name == style)
    ]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected one paragraph starting with {prefix!r}; found {len(matches)}"
        )
    return matches[0]


def set_paragraph(paragraph: Paragraph, text: str) -> None:
    run_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        run_properties = deepcopy(paragraph.runs[0]._r.rPr)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if run_properties is not None:
        run._r.insert(0, run_properties)


def set_by_prefix(doc: Document, prefix: str, text: str, style: str | None = None) -> None:
    set_paragraph(find_one(doc, prefix, style), text)


def insert_after(anchor: Paragraph, text: str, style: str = "Normal") -> Paragraph:
    element = OxmlElement("w:p")
    anchor._p.addnext(element)
    paragraph = Paragraph(element, anchor._parent)
    paragraph.style = style
    paragraph.add_run(text)
    return paragraph


def remove_empty_between(start, end) -> None:
    element = start.getnext()
    while element is not None and element is not end:
        next_element = element.getnext()
        content = [child for child in element if child.tag != qn("w:pPr")]
        if element.tag == qn("w:p") and not content:
            element.getparent().remove(element)
        element = next_element


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    run_properties = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        run_properties = deepcopy(paragraph.runs[0]._r.rPr)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(str(text))
    if run_properties is not None:
        run._r.insert(0, run_properties)
    for extra in list(cell.paragraphs[1:]):
        cell._tc.remove(extra._p)


def set_fill(cell, fill: str | None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    if fill is None:
        shading.attrib.pop(qn("w:fill"), None)
    else:
        shading.set(qn("w:fill"), fill)


def refresh_best_cells(table, best_cells: list[tuple[int, int]]) -> None:
    for row in table.rows:
        for cell in row.cells:
            shading = cell._tc.tcPr.find(qn("w:shd"))
            if shading is not None and shading.get(qn("w:fill")) == "E2EFDA":
                set_fill(cell, None)
    for row_index, column_index in best_cells:
        set_fill(table.rows[row_index].cells[column_index], "E2EFDA")


def replace_picture_before_caption(
    doc: Document,
    caption_prefix: str,
    image: Path,
    scale: float = 1.0,
) -> None:
    captions = [
        paragraph
        for paragraph in doc.paragraphs
        if normalized(paragraph.text).startswith(normalized(caption_prefix))
        and paragraph.style.name == "Caption"
    ]
    caption = captions[0] if len(captions) == 1 else find_one(doc, caption_prefix)
    element = caption._p.getprevious()
    while element is not None and not element.xpath(".//a:blip"):
        element = element.getprevious()
    if element is None:
        raise RuntimeError(f"No picture found before {caption_prefix!r}")
    extent = element.xpath(".//wp:extent")
    if not extent:
        raise RuntimeError(f"Picture extent missing before {caption_prefix!r}")
    width = Emu(round(int(extent[0].get("cx")) * scale))
    height = Emu(round(int(extent[0].get("cy")) * scale))
    paragraph = Paragraph(element, caption._parent)
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image), width=width, height=height)
    remove_empty_between(paragraph._p, caption._p)
    paragraph.paragraph_format.keep_with_next = True


def row_for(frame: pd.DataFrame, **filters) -> pd.Series:
    result = frame
    for column, value in filters.items():
        result = result[result[column] == value]
    if len(result) != 1:
        raise RuntimeError(f"Expected one row for {filters}; found {len(result)}")
    return result.iloc[0]


def update_summary_and_abstract(doc: Document) -> None:
    set_by_prefix(
        doc,
        "Esta tesis evalúa si la incorporación",
        "Esta tesis evalúa si la incorporación de variables exógenas —el índice de "
        "volatilidad implícita (VIX) y un indicador de tono de noticias construido "
        "con GDELT— mejora el pronóstico de los retornos logarítmicos acumulados "
        "del índice NASDAQ-100 en horizontes de 1, 5, 10 y 20 días hábiles. Se "
        "comparan RW, ARIMA, ARIMAX/SARIMAX y tres redes LSTM. Tras una auditoría "
        "de procedencia, los datos se reconstruyeron desde fuentes identificables: "
        "cierre diario NASDAQ-100 de FRED/Nasdaq, cierre del VIX de Cboe y un proxy "
        "de tono del mercado Nasdaq/tecnología derivado de GDELT. El conjunto "
        "canónico contiene 2.558 observaciones entre el 19-02-2015 y el 22-04-2025. "
        "El protocolo combina OOS 2024, validación walk-forward mensual, robustez "
        "enero–abril de 2025, pruebas Diebold–Mariano con HAC y corrección "
        "Harvey–Leybourne–Newbold, calibración Mincer–Zarnowitz, acierto direccional "
        "Pesaran–Timmermann y un análisis de diez semillas por variante LSTM.",
    )
    set_by_prefix(
        doc,
        "Los resultados favorecen de manera uniforme",
        "En OOS 2024, el RW obtuvo RMSE 0.01146 en T+1; ARIMAX/SARIMAX fueron "
        "significativamente peores (p_HLN = 0.0373). En T+20, LSTM+Sent+VIX "
        "registró el menor RMSE (0.03983 frente a 0.04368 del RW), pero sin "
        "diferencia significativa (p_HLN = 0.1811). La LSTM univariada sí produjo "
        "un contraste favorable en la semilla canónica (p_HLN = 0.0201), aunque "
        "su mediana en diez semillas fue 0.0562, cinco semillas quedaron bajo "
        "0.05 y el ensemble obtuvo p = 0.0584, por lo que la inferencia depende de "
        "la inicialización. En walk-forward, LSTM+Sent+VIX alcanzó el menor RMSE "
        "medio en T+20 (0.03593 frente a 0.04136 del RW). En 2025, el RW volvió a "
        "ser el mejor modelo en T+5, T+10 y T+20, sin diferencias significativas "
        "contra él bajo HLN. Los bloques walk-forward muestran reversión en los "
        "bloques 3, 7 y 12, mientras la estratificación por VIX observado en el "
        "origen favorece a los modelos en los cuartiles altos; ambas lecturas no "
        "son equivalentes porque el VIX de origen no anticipa cambios dentro de la "
        "ventana futura. Se concluye que existen mejoras puntuales de error, pero "
        "no una superioridad estable entre horizontes, periodos e inicializaciones.",
    )
    set_by_prefix(
        doc,
        "This thesis evaluates whether exogenous information",
        "This thesis evaluates whether exogenous information — the implied-volatility "
        "index (VIX) and a GDELT-based news-tone indicator — improves forecasts of "
        "accumulated log returns of the NASDAQ-100 at 1, 5, 10, and 20 trading-day "
        "horizons. RW, ARIMA, ARIMAX/SARIMAX, and three LSTM variants are compared. "
        "Following a provenance audit, the data were rebuilt from identifiable "
        "sources: the daily NASDAQ-100 close from FRED/Nasdaq, the VIX close from "
        "Cboe, and a GDELT-derived proxy for Nasdaq/technology-market news tone. "
        "The canonical dataset contains 2,558 observations from 19 February 2015 "
        "to 22 April 2025. The protocol combines a 2024 out-of-sample evaluation, "
        "monthly walk-forward validation, a January–April 2025 robustness block, "
        "Diebold–Mariano tests with HAC errors and the Harvey–Leybourne–Newbold "
        "correction, Mincer–Zarnowitz calibration, Pesaran–Timmermann directional "
        "testing, and a ten-seed analysis for each LSTM variant.",
    )
    set_by_prefix(
        doc,
        "Results uniformly favor weak-form market efficiency",
        "In the 2024 out-of-sample evaluation, RW achieved RMSE 0.01146 at T+1; "
        "ARIMAX/SARIMAX were significantly worse (p_HLN = 0.0373). At T+20, "
        "LSTM+Sent+VIX produced the lowest RMSE (0.03983 versus 0.04368 for RW), "
        "without a significant difference (p_HLN = 0.1811). The univariate LSTM "
        "was significant under the canonical seed (p_HLN = 0.0201), but its "
        "ten-seed median was 0.0562, five seeds fell below 0.05, and the ensemble "
        "yielded p = 0.0584, showing initialization-sensitive inference. In "
        "walk-forward validation, LSTM+Sent+VIX achieved the lowest mean T+20 RMSE "
        "(0.03593 versus 0.04136 for RW). In 2025, RW was again best at T+5, T+10, "
        "and T+20, with no HLN-significant comparison against it. Walk-forward "
        "blocks 3, 7, and 12 show reversals, whereas stratification by VIX observed "
        "at forecast origin favors the models in high-VIX quartiles; these views "
        "are not equivalent because origin-date VIX does not anticipate changes "
        "inside the future forecast window. The evidence supports pointwise error "
        "improvements, but not stable superiority across horizons, periods, and "
        "initializations.",
    )


def update_methodology(doc: Document) -> None:
    anchor = find_one(doc, "Datos y variables", "Heading 2")
    additions = [
        ("Procedencia y reconstrucción de las fuentes", "Heading 3"),
        (
            "El repositorio original de adquisición se perdió antes de consolidar "
            "esta versión. El CSV integrado heredado se conserva bajo SHA-256, pero "
            "no permite demostrar el origen de cada campo. La auditoría forense "
            "determinó que su objetivo correspondía al cierre ajustado de QQQ y no "
            "al nivel del NASDAQ-100; que VIX_Close reproducía la apertura del VIX "
            "por un error de encabezados, con un último dato rellenado; y que el "
            "sentimiento tenía origen no verificable y permanecía constante en las "
            "últimas 75 filas. Ese archivo se usa únicamente como evidencia histórica "
            "y no alimenta los resultados reportados.",
            "Normal",
        ),
        (
            "La reconstrucción usa el cierre diario de la serie NASDAQ100 publicada "
            "por FRED con fuente original Nasdaq [30], el cierre diario oficial del "
            "VIX publicado por Cboe [31] y GDELT 2.0 GKG mediante el conjunto público "
            "gdelt-bq.gdeltv2.gkg_partitioned [32]. La descarga y la consulta quedan "
            "registradas mediante URL, fecha, parámetros, job de BigQuery, consulta "
            "SQL y SHA-256 en data/manifests/.",
            "Normal",
        ),
        (
            "El sentimiento se define como la media diaria, ponderada por número de "
            "artículos, del primer componente de V2Tone. El alcance sellado "
            "nasdaq_market exige una mención organizacional a Nasdaq y el tema "
            "ECON_STOCKMARKET. Se eligió por cobertura y relevancia documental antes "
            "de ejecutar los modelos, nunca por desempeño predictivo. Por ello se "
            "interpreta como proxy del tono de noticias del mercado Nasdaq/tecnología, "
            "no como un corpus exclusivo de las empresas del índice.",
            "Normal",
        ),
        (
            "El calendario canónico es la intersección de fechas observadas del "
            "NASDAQ-100 y el VIX. Las noticias de fines de semana se asignan a la "
            "primera fecha de mercado posterior. La fecha 2019-04-19 se excluye por "
            "no tener observación VIX y 2017-08-29 por ausencia de tono GDELT; no se "
            "aplica forward-fill, interpolación ni imputación neutral. El conjunto "
            "resultante contiene 2.558 filas entre 2015-02-19 y 2025-04-22, con "
            "SHA-256 abf82d900314ce09cd00113785804e9f6742c38dd0867880ae9620c2a12f92ce.",
            "Normal",
        ),
        ("Definición de variables", "Heading 3"),
    ]
    for text, style in additions:
        anchor = insert_after(anchor, text, style)

    set_by_prefix(
        doc,
        "En los datos del estudio se obtuvo p = 0.302",
        "En los datos reconstruidos se obtuvo p = 0.4698, por lo que los términos "
        "estacionales permanecieron desactivados en todos los horizontes.",
    )
    set_by_prefix(
        doc,
        "Scripts: src/stages/00_prepare.py",
        "Scripts: scripts/data/ contiene la adquisición, auditoría e integración de "
        "fuentes; src/stages/00_prepare.py genera features y splits; "
        "src/stages/05_stationarity.py produce ADF/KPSS y correlogramas; las etapas "
        "10_baselines.py–16_robustez_2025.py entrenan y validan; 20_evaluate_stats.py "
        "y 24_wf_metrics.py calculan métricas e inferencia; 30_make_figures.py y "
        "31_regimes.py generan figuras y análisis por régimen.",
    )
    set_by_prefix(
        doc,
        "Pruebas: reports/data/tbl_DM_OOS.csv",
        "Pruebas: reports/data/tbl_DM_OOS.csv, tbl_MZ_core.csv, tbl_PT_OOS.csv, "
        "tbl_DM_2025.csv, multiseed_dm_T20_all_variants.csv y "
        "stationarity_tests.csv.",
    )


def update_tables(doc: Document) -> None:
    metrics = pd.read_csv(ROOT / "reports/data/metrics_OOS.csv")
    wf = pd.read_csv(ROOT / "reports/data/metrics_WF.csv")
    dm = pd.read_csv(ROOT / "reports/data/tbl_DM_OOS.csv")
    multiseed = pd.read_csv(ROOT / "reports/data/multiseed_summary.csv")
    robust = pd.read_csv(ROOT / "reports/data/metrics_OOS_2025.csv")
    stationarity = pd.read_csv(ROOT / "reports/data/stationarity_tests.csv")

    table = doc.tables[5]
    for row_index, model in enumerate(MODEL_ORDER, start=2):
        set_cell_text(table.rows[row_index].cells[0], MODEL_LABELS[model])
        for horizon_index, horizon in enumerate((1, 5, 10, 20)):
            record = row_for(metrics, Horizon=horizon, Model=model)
            start = 1 + horizon_index * 3
            set_cell_text(table.rows[row_index].cells[start], f"{record.RMSE:.5f}")
            set_cell_text(table.rows[row_index].cells[start + 1], f"{record.MAE:.5f}")
            mda = "n.a." if model == "RW" else f"{record.MDA:.3f}"
            set_cell_text(table.rows[row_index].cells[start + 2], mda)
    refresh_best_cells(table, [(2, 1), (3, 4), (7, 7), (7, 10)])

    table = doc.tables[6]
    for row_index, model in enumerate(MODEL_ORDER[1:], start=1):
        set_cell_text(table.rows[row_index].cells[0], MODEL_LABELS[model])
        for column_index, horizon in enumerate((1, 5, 10, 20), start=1):
            record = row_for(dm, Horizon=horizon, Challenger=model, Benchmark="RW")
            set_cell_text(
                table.rows[row_index].cells[column_index],
                f"{record.DM_HLN:.2f} (p={record.p_HLN:.3f})",
            )

    table = doc.tables[7]
    for row_index, model in enumerate(MODEL_ORDER, start=2):
        set_cell_text(table.rows[row_index].cells[0], MODEL_LABELS[model])
        for horizon_index, horizon in enumerate((1, 5, 10, 20)):
            record = row_for(wf, Horizon=horizon, Model=model)
            start = 1 + horizon_index * 2
            set_cell_text(table.rows[row_index].cells[start], f"{record.RMSE:.5f}")
            mda = "n.a." if model == "RW" else f"{record.MDA:.3f}"
            set_cell_text(table.rows[row_index].cells[start + 1], mda)
    refresh_best_cells(table, [(5, 1), (3, 3), (7, 5), (7, 7)])

    table = doc.tables[8]
    for row_index, model in enumerate(LSTM_ORDER, start=1):
        set_cell_text(table.rows[row_index].cells[0], MODEL_LABELS[model])
        for column_index, horizon in enumerate((1, 5, 10, 20), start=1):
            record = row_for(multiseed, Variant=model, Horizon=horizon)
            set_cell_text(
                table.rows[row_index].cells[column_index],
                f"{record.RMSE_mean:.5f} ± {record.RMSE_std:.5f}",
            )

    table = doc.tables[9]
    for row_index, model in enumerate(MODEL_ORDER, start=2):
        set_cell_text(table.rows[row_index].cells[0], MODEL_LABELS[model])
        for horizon_index, horizon in enumerate((1, 5, 10, 20)):
            record = row_for(robust, Horizon=horizon, Model=model)
            start = 1 + horizon_index * 2
            set_cell_text(table.rows[row_index].cells[start], f"{record.RMSE:.5f}")
            mda = "n.a." if model == "RW" else f"{record.MDA:.3f}"
            set_cell_text(table.rows[row_index].cells[start + 1], mda)
    refresh_best_cells(table, [(4, 1), (2, 3), (2, 5), (2, 7)])

    table = doc.tables[11]
    for row_index, record in enumerate(stationarity.itertuples(), start=1):
        values = [
            record.series,
            record.transform,
            record.test,
            str(int(record.lag)),
            f"{record.stat:.2f}",
            f"{record.pvalue:.2f}",
            record.conclusion,
        ]
        for column_index, value in enumerate(values):
            set_cell_text(table.rows[row_index].cells[column_index], value)


def update_results(doc: Document) -> None:
    replacements = [
        (
            "Mejor RMSE por horizonte sombreado en verde. ¹",
            "Mejor RMSE por horizonte sombreado en verde. ¹ La regla de activación "
            "estacional (Kruskal–Wallis, p = 0.4698 > 0.10) no activó m = 5, por "
            "lo que SARIMAX coincide con ARIMAX. ² El RW pronostica retorno acumulado "
            "nulo; su MDA no es aplicable.",
        ),
        (
            "La Tabla 7.1 presenta las métricas",
            "La Tabla 7.1 presenta las métricas OOS 2024. En T+1, RW obtuvo RMSE "
            "0.01146; LSTM y LSTM+Sent+VIX quedaron prácticamente empatadas "
            "(0.01147), mientras ARIMAX/SARIMAX y LSTM+Sent aumentaron el error a "
            "0.01172 y 0.01189. La estrechez de las diferencias no implica igualdad "
            "inferencial: la Sección 7.2 muestra que ARIMAX/SARIMAX son peores que "
            "RW en este horizonte bajo DM-HLN.",
        ),
        (
            "A medida que el horizonte se amplía",
            "En T+5, ARIMA registró el menor RMSE (0.02418); en T+10 y T+20 el "
            "liderazgo correspondió a LSTM+Sent+VIX (0.03302 y 0.03983). En T+20, "
            "los seis modelos presentaron RMSE inferior al RW (0.04368), con una "
            "reducción puntual de 8.8% para LSTM+Sent+VIX. El MDA de los modelos "
            "crece hasta 0.70–0.73 en T+20, pero el solapamiento de retornos y los "
            "casos de predicciones de un solo signo impiden interpretarlo por sí "
            "solo como capacidad direccional.",
        ),
        (
            "Errores HAC con kernel de Bartlett",
            "Errores HAC con kernel de Bartlett (L = h−1) y corrección HLN. Dos "
            "contrastes sustantivos frente a RW alcanzan 5%: ARIMAX/SARIMAX es peor "
            "en T+1 (p = 0.037) y LSTM es mejor en T+20 para la semilla canónica "
            "(p = 0.020). ARIMAX y SARIMAX son predicciones idénticas y no deben "
            "contarse como evidencias independientes.",
        ),
        (
            "La Tabla 7.2 reporta la prueba",
            "La Tabla 7.2 reporta DM frente al RW con HAC y corrección HLN. En T+1, "
            "ARIMAX/SARIMAX produce DM_HLN = −2.09 (p = 0.0373), evidencia de mayor "
            "pérdida que RW. En T+20, la LSTM univariada produce DM_HLN = 2.34 "
            "(p = 0.0201), mientras ARIMA, ARIMAX/SARIMAX, LSTM+Sent y "
            "LSTM+Sent+VIX no rechazan igualdad predictiva (p entre 0.0819 y "
            "0.2232). Dado que se realizan múltiples contrastes y no se aplica una "
            "corrección familiar, ambos resultados se interpretan como evidencia "
            "puntual. La robustez de la LSTM a la semilla se evalúa en la Sección 7.5.",
        ),
        (
            "La prueba de Pesaran–Timmermann",
            "La prueba de Pesaran–Timmermann [28] no detecta habilidad direccional "
            "significativa: el menor p-valor calculable es 0.4155. Varias celdas de "
            "horizontes largos no son calculables porque el modelo pronostica un "
            "solo signo; en esos casos la prueba carece de varianza. Por tanto, los "
            "MDA altos de T+20 no aportan evidencia inferencial independiente.",
        ),
        (
            "Una nota de transparencia metodológica",
            "Una nota de transparencia metodológica es obligada. Las auditorías "
            "del proyecto detectaron, antes de esta reproducción, una fuga por "
            "objetivos solapados y un desplazamiento de 40 filas en la verificación "
            "de y_true; ambos invalidaron resultados intermedios. La auditoría de "
            "procedencia posterior mostró además que el CSV heredado no representaba "
            "las fuentes declaradas. Las tablas actuales no corrigen cifras antiguas "
            "de forma selectiva: proceden de una ejecución completa sobre el dataset "
            "reconstruido, después de cerrar los tres problemas. Esta nota documenta "
            "el proceso de control y no constituye evidencia inferencial adicional.",
        ),
        (
            "Las regresiones de Mincer–Zarnowitz",
            "Las regresiones de Mincer–Zarnowitz en niveles muestran que, en T+1, "
            "la mayoría de modelos no rechaza α = 0 ni β = 1 al 5%; la excepción es "
            "LSTM+Sent (p_α = 0.0054; p_β=1 = 0.0035), mientras ARIMAX/SARIMAX "
            "queda en el límite (p_β=1 = 0.0600). Desde T+5, todos los modelos "
            "rechazan ambos contrastes por separado. En T+20, las pendientes se "
            "sitúan entre 0.7808 y 0.8153 y R² entre 0.6998 y 0.7200. La pérdida de "
            "calibración es, por tanto, general en horizontes largos, pero no uniforme "
            "en T+1.",
        ),
        (
            "Mejor RMSE por horizonte sombreado en verde. En los promedios",
            "Mejor RMSE por horizonte sombreado en verde. En T+1 solo las tres "
            "variantes LSTM mejoran el promedio de RW; en T≥5 todos los modelos "
            "quedan por debajo. Los bloques 3, 7 y 12 muestran reversión simultánea "
            "en T+20, pero la estratificación contemporánea por VIX no reproduce una "
            "división simple entre calma y estrés.",
        ),
        (
            "La Tabla 7.3 resume el promedio",
            "La Tabla 7.3 resume doce bloques mensuales walk-forward con re-estimación "
            "por bloque. Los mejores RMSE son LSTM en T+1 (0.01100), ARIMA en T+5 "
            "(0.02276) y LSTM+Sent+VIX en T+10 y T+20 (0.03004 y 0.03593). En T+20, "
            "todos los modelos mejoran el RW de 0.04136; el orden relativo coincide "
            "parcialmente con OOS, donde LSTM+Sent+VIX también registra el menor error.",
        ),
        (
            "La heterogeneidad entre bloques",
            "La heterogeneidad entre bloques matiza los promedios. En T+20, todas "
            "las familias empeoran frente al RW en los bloques 3, 7 y 12, mientras "
            "la mayoría de los bloques restantes favorece a los modelos. Sin embargo, "
            "los cuartiles definidos por VIX observado en la fecha de origen ofrecen "
            "otra lectura: en Q1 solo LSTM mejora marginalmente al RW (0.06038 frente "
            "a 0.06062), mientras en Q3 y Q4 las menores cifras corresponden a "
            "LSTM+Sent+VIX (0.02869) y ARIMAX/SARIMAX (0.02536), frente a RW de "
            "0.03509 y 0.03979. No hay contradicción mecánica: el régimen de origen "
            "describe información contemporánea y no identifica cambios que ocurren "
            "dentro de los 20 días posteriores. Estos resultados son descriptivos y "
            "no prueban una regla predictiva de conmutación por régimen.",
        ),
        (
            "Para LSTM+Sent+VIX en T = 20 frente a RW",
            "En T+20, LSTM presenta mediana p = 0.0562, cinco de diez semillas bajo "
            "0.05 y ensemble p = 0.0584. LSTM+Sent presenta mediana p = 0.1893, "
            "cero semillas bajo 0.05 y ensemble p = 0.1523. LSTM+Sent+VIX obtiene "
            "el menor RMSE medio (0.04033 ± 0.00070), pero mediana p = 0.2225, una "
            "semilla bajo 0.05 y ensemble p = 0.2101.",
        ),
        (
            "Los resultados neuronales presentados hasta aquí",
            "Los resultados neuronales canónicos usan semilla 42. Para medir la "
            "dependencia de esa elección, cada una de las tres variantes se reentrenó "
            "con diez semillas en los cuatro horizontes, para 120 entrenamientos. La "
            "Tabla 7.4 resume media y desviación del RMSE; la prueba DM-HLN por semilla "
            "y por ensemble se conserva en multiseed_dm_T20_all_variants.csv.",
        ),
        (
            "Tres regularidades emergen",
            "La inferencia cambia materialmente con la inicialización. La LSTM "
            "canónica es significativa en T+20 (p = 0.0201), pero la mediana de diez "
            "semillas queda en 0.0562 y el ensemble en 0.0584. LSTM+Sent no presenta "
            "ninguna semilla significativa. LSTM+Sent+VIX mantiene el mejor RMSE "
            "medio entre variantes, pero solo una semilla es significativa y su "
            "ensemble no lo es. Por ello, el signo favorable del error en T+20 no "
            "autoriza una afirmación de superioridad robusta; la semilla debe tratarse "
            "como dimensión del análisis, no como detalle de implementación.",
        ),
        (
            "Mejor RMSE por horizonte sombreado en verde. El periodo incluye",
            "Mejor RMSE por horizonte sombreado en verde. En T+1, ARIMAX/SARIMAX "
            "presenta el menor RMSE (0.02186). En T+5, T+10 y T+20, el RW presenta "
            "el menor error. Ningún contraste DM-HLN contra RW alcanza 5%.",
        ),
        (
            "Como prueba de estrés",
            "Como prueba de estrés, los modelos se evaluaron entre enero y abril de "
            "2025. Las LSTM se entrenaron con información hasta diciembre de 2024 y "
            "los modelos clásicos continuaron con ventana expanding y re-estimación "
            "mensual. El RMSE de T+1 aumentó aproximadamente de 0.011–0.012 en 2024 "
            "a 0.022 en este bloque.",
        ),
        (
            "El resultado central del bloque es categórico",
            "En T+5, T+10 y T+20, RW registra los menores RMSE (0.04087, 0.05277 y "
            "0.07147). En T+20, los modelos quedan entre 12.6% y 27.0% por encima del "
            "RW. Ninguna comparación contra RW es significativa con HLN. La única "
            "celda significativa de tbl_DM_2025.csv compara LSTM con SARIMAX en "
            "T+20 (p_HLN = 0.0496), no con RW, y se interpreta como hallazgo aislado. "
            "El MDA de todos los modelos en T+20 es 0.2364 porque sus acumulados "
            "pronosticados tienen el mismo signo durante gran parte de la caída.",
        ),
        (
            "Este bloque cierra el arco",
            "El bloque de 2025 muestra que las mejoras de error observadas en OOS y "
            "walk-forward 2024 no se trasladan de manera estable a otro periodo. "
            "También impide convertir la estratificación contemporánea por VIX de "
            "2024 en una regla operacional: conocer el cuartil en el origen no "
            "equivale a anticipar un cambio de régimen dentro del horizonte.",
        ),
        (
            "Tomados en conjunto, los resultados",
            "Tomados en conjunto, los resultados no sostienen un modelo dominante. "
            "En T+1, ARIMAX/SARIMAX es significativamente peor que RW; en T+20, la "
            "LSTM canónica es significativamente mejor, pero esa inferencia se debilita "
            "al variar la semilla. LSTM+Sent+VIX obtiene el menor RMSE en T+10 y T+20 "
            "OOS y walk-forward, sin diferencia DM-HLN frente a RW y sin conservar el "
            "liderazgo en 2025. La evidencia, por tanto, es compatible con mejoras "
            "puntuales, no con superioridad estable y generalizable.",
        ),
        (
            "La segunda lectura concierne",
            "La segunda lectura concierne a complejidad y variables exógenas. "
            "LSTM+Sent+VIX alcanza los mejores errores puntuales de horizonte largo, "
            "pero el sentimiento por sí solo no domina a la LSTM univariada y ninguna "
            "variante exógena demuestra una ventaja inferencial robusta. ARIMA y "
            "ARIMAX/SARIMAX permanecen competitivos con menor complejidad, especialmente "
            "en T+5 y en varios bloques walk-forward.",
        ),
        (
            "La tercera lectura concierne",
            "La tercera lectura concierne a la estabilidad temporal. Los bloques "
            "3, 7 y 12 revierten simultáneamente las ventajas en T+20, pero los "
            "cuartiles de VIX en el origen favorecen a varios modelos en Q3 y Q4. "
            "La diferencia revela que una etiqueta contemporánea no describe el "
            "régimen futuro de toda la ventana. Se requiere una prueba prospectiva "
            "específica antes de afirmar que el desempeño depende de calma o estrés.",
        ),
        (
            "La cuarta lectura es metodológica",
            "La cuarta lectura es metodológica. La reconstrucción de fuentes mostró "
            "que un pipeline puede ser computacionalmente reproducible y, aun así, "
            "carecer de procedencia demostrable. Los contratos anti-fuga, la "
            "verificación de y_true, los hashes de fuentes y artefactos, las corridas "
            "limpias y el análisis multi-semilla cubren riesgos distintos y "
            "complementarios. En esta versión, dos ejecuciones independientes "
            "reprodujeron 512/512 predicciones y todas las tablas comparables byte a "
            "byte; las figuras también quedaron deterministas tras fijar el jitter.",
        ),
    ]
    for prefix, text in replacements:
        set_by_prefix(doc, prefix, text)

    captions = {
        "Figura 7.2": "Figura 7.2 — Diagrama volcano de Diebold–Mariano frente al RW en T+20 (OOS 2024): diferencia de RMSE (eje X) y −log₁₀(p_HLN) con HAC L=h−1 (eje Y). Las líneas marcan p=0.05 y p=0.10.",
        "Figura 7.4": "Figura 7.4 — Mapa de calor walk-forward en T+20: diferencia de RMSE frente al RW (%) por modelo y bloque mensual de 2024. Azul indica ventaja sobre RW; rojo, desventaja. Los bloques 3, 7 y 12 muestran reversión simultánea.",
        "Figura 7.5": "Figura 7.5 — Diferencia acumulada de pérdida cuadrática frente al RW en T+20 durante 2024. Valores bajo cero representan menor pérdida acumulada del modelo.",
        "Figura 7.6": "Figura 7.6 — Distribución del RMSE OOS de las variantes LSTM sobre 10 semillas por horizonte; la línea discontinua representa el RW. La dispersión cuantifica sensibilidad a la inicialización.",
    }
    for prefix, text in captions.items():
        set_by_prefix(doc, prefix, text, "Caption")


def update_conclusions(doc: Document) -> None:
    replacements = [
        (
            "Esta tesis se propuso determinar",
            "Esta tesis evaluó si VIX y tono de noticias GDELT mejoran el pronóstico "
            "multi-horizonte del NASDAQ-100 bajo un protocolo reproducible. La "
            "respuesta es matizada: hay reducciones puntuales de RMSE y dos contrastes "
            "DM-HLN frente a RW al 5%, uno desfavorable a ARIMAX/SARIMAX en T+1 y "
            "otro favorable a la LSTM canónica en T+20. Sin embargo, el segundo es "
            "sensible a la semilla, las ventajas no se mantienen en 2025 y ningún "
            "modelo domina de forma estable entre horizontes y periodos.",
        ),
        (
            "Se construyó un pipeline íntegramente reproducible",
            "Se reconstruyó el conjunto de datos desde FRED/Nasdaq, Cboe y GDELT y "
            "se versionaron URLs, parámetros, consulta SQL, job y hashes. El ejercicio "
            "también documentó que el archivo heredado no era una fuente primaria: "
            "su objetivo era QQQ ajustado, VIX_Close correspondía a la apertura y el "
            "sentimiento no tenía procedencia verificable. Sobre el conjunto nuevo, "
            "el pipeline completo terminó con RUN_ALL_OK; dos corridas independientes "
            "reprodujeron las 512 predicciones y las tablas comparables.",
        ),
        (
            "Los siete modelos efectivos se implementaron",
            "Los siete modelos efectivos se ejecutaron sobre retornos diarios con "
            "pronóstico iterado, exógenas rezagadas y congeladas, y re-estimación "
            "según el protocolo. El gate Kruskal–Wallis obtuvo p = 0.4698, por lo que "
            "m = 5 permaneció desactivado y SARIMAX coincidió con ARIMAX. ADF y KPSS "
            "sobre el objetivo oficial confirman no estacionariedad en nivel y "
            "estacionariedad tras una diferencia.",
        ),
        (
            "En el horizonte diario",
            "En T+1, RW obtuvo RMSE 0.01146 y ARIMAX/SARIMAX fue significativamente "
            "peor (p_HLN = 0.0373). En T+20, LSTM+Sent+VIX logró el menor RMSE OOS "
            "(0.03983) y walk-forward (0.03593), pero su contraste OOS frente a RW "
            "no fue significativo (p_HLN = 0.1811). La LSTM univariada sí fue "
            "significativa con semilla 42 (p_HLN = 0.0201), sin respaldo estable al "
            "variar la inicialización. Pesaran–Timmermann no detectó habilidad "
            "direccional y Mincer–Zarnowitz mostró descalibración general desde T+5.",
        ),
        (
            "Las tres pruebas de robustez convergen",
            "Las pruebas de robustez delimitan el alcance. Para la LSTM en T+20, "
            "cinco de diez semillas quedaron bajo 0.05, pero la mediana fue 0.0562 "
            "y el ensemble 0.0584. LSTM+Sent+VIX obtuvo el menor RMSE medio, con una "
            "sola semilla significativa y ensemble p = 0.2101. En 2025, RW fue el "
            "mejor modelo en T+5, T+10 y T+20 y ninguna comparación contra RW fue "
            "significativa bajo HLN. Los regímenes de VIX en el origen y las "
            "reversiones por bloque describen aspectos diferentes y no justifican "
            "una regla de conmutación sin validación prospectiva.",
        ),
        (
            "Para un practicante",
            "Para un practicante, las cifras no justifican seleccionar un modelo por "
            "una sola corrida o un único periodo. La menor pérdida puntual de "
            "LSTM+Sent+VIX en horizontes largos debe ponderarse con su falta de "
            "significancia frente a RW, la sensibilidad a semilla y el deterioro en "
            "2025. Estos resultados evalúan exactitud estadística, no rentabilidad, "
            "y no constituyen una recomendación de inversión.",
        ),
        (
            "La experiencia de este proyecto sugiere",
            "La experiencia del proyecto muestra que reproducibilidad computacional "
            "y procedencia son requisitos distintos. Un resultado puede repetirse "
            "desde un CSV y seguir sin ser auditable si no se preservan fuentes, "
            "campos, calendario y transformaciones. Del mismo modo, una semilla "
            "canónica puede producir una conclusión que no representa la distribución "
            "de inicializaciones. Ambos riesgos deben reportarse explícitamente.",
        ),
        (
            "La contribución que este trabajo considera",
            "La contribución metodológica es un protocolo que enlaza fuente, dataset, "
            "predicción, métrica, tabla y figura mediante manifiestos y pruebas. La "
            "reconstrucción no se usó para preservar resultados anteriores, sino para "
            "volver a ejecutar todo el estudio. Esta separación evita presentar como "
            "continuidad empírica lo que en realidad es una nueva base de evidencia.",
        ),
        (
            "Los alcances de las conclusiones",
            "Los alcances están acotados por siete limitaciones. Primero, el repositorio "
            "original de adquisición se perdió y la procedencia del sentimiento "
            "heredado sigue sin poder recuperarse; por ello los resultados actuales "
            "pertenecen a una reconstrucción y no validan las cifras históricas. "
            "Segundo, nasdaq_market es un proxy de noticias Nasdaq/tecnología, no un "
            "corpus exclusivo del índice, y el primer componente de V2Tone no fue "
            "validado manualmente como medida semántica. Tercero, los archivos fuente "
            "no se redistribuyen en Git mientras se revisan licencias; la reproducción "
            "desde cero depende de proveedores externos, aunque hashes y consultas "
            "quedan versionados. Cuarto, se estudia un solo índice y un periodo OOS "
            "principal, con un bloque 2025 breve. Quinto, los cuartiles de VIX son "
            "contemporáneos al origen y no anticipan el régimen futuro. Sexto, las "
            "exógenas se congelan y la arquitectura LSTM no agota el espacio de "
            "modelos. Séptimo, no se modelan costos ni reglas de negociación.",
        ),
        (
            "Las limitaciones trazan la agenda",
            "El trabajo futuro debe comenzar por preservación: depositar una copia "
            "permitida de las fuentes y el dataset canónico en almacenamiento "
            "institucional con control de acceso, manteniendo los hashes públicos. "
            "Después, conviene replicar el protocolo en otros índices y periodos, "
            "evaluar alcances GDELT preespecificados sin selección por rendimiento, "
            "modelar dinámicamente las exógenas y probar reglas de régimen de forma "
            "prospectiva. También se recomienda incorporar inferencia para múltiples "
            "comparaciones y mantener el análisis multi-semilla como requisito de "
            "cualquier extensión neuronal.",
        ),
    ]
    for prefix, text in replacements:
        set_by_prefix(doc, prefix, text)


def update_appendices(doc: Document) -> None:
    set_by_prefix(
        doc,
        "Muestras: mismas ventanas",
        "Muestra: conjunto canónico completo de 2.558 observaciones entre "
        "2015-02-19 y 2025-04-22. El diagnóstico es descriptivo de la transformación "
        "y no se usa para seleccionar el modelo ganador.",
    )
    set_by_prefix(
        doc,
        "Especificación ADF:",
        "Especificación ADF: constante, sin tendencia determinista, rezagos fijos "
        "0, 1 y 2 y autolag desactivado.",
    )
    set_by_prefix(
        doc,
        "Especificación KPSS:",
        "Especificación KPSS: estacionariedad alrededor de constante y rezagos de "
        "varianza de largo plazo fijados en 0, 1 y 2.",
    )
    set_by_prefix(
        doc,
        "Los resultados muestran no estacionariedad",
        "Sobre el NASDAQ-100 oficial, ADF no rechaza raíz unitaria en nivel "
        "(p = 0.885–0.905) y KPSS rechaza estacionariedad (p = 0.01). En primera "
        "diferencia, ADF rechaza raíz unitaria (p reportado 0.00) y KPSS no rechaza "
        "estacionariedad (p = 0.10). Esto respalda el uso de ret_1d y el orden "
        "(1,0,1) sobre retornos, equivalente a una diferencia sobre log(P).",
    )

    set_by_prefix(
        doc,
        "Este apéndice documenta el procedimiento verificado",
        "Este apéndice documenta la reconstrucción de datos y la reproducción "
        "verificada de los resultados sellados el 16 de julio de 2026 en Windows 11 "
        "con Python 3.12.13. El dataset, los árboles de artefactos y los resultados "
        "clave están fijados en data/manifests/reproduction_results.json.",
    )
    set_by_prefix(
        doc,
        "El proyecto se ejecuta con Python 3.11",
        "El proyecto se ejecuta con Python 3.12 y dependencias fijadas en "
        "requirements.txt; la reconstrucción de fuentes añade requirements-data.txt.",
    )
    set_by_prefix(
        doc,
        "git clone --branch correcciones-auditoria-final",
        "git clone --branch reconstruccion-datos-origen https://github.com/joslat92/TesisProject.git\n"
        "cd TesisProject\n"
        "python -m venv .venv\n"
        ".venv\\Scripts\\python.exe -m pip install --upgrade pip\n"
        ".venv\\Scripts\\python.exe -m pip install -r requirements.txt\n"
        ".venv\\Scripts\\python.exe -m pip install -r requirements-data.txt",
    )

    install = find_one(doc, "git clone --branch reconstruccion-datos-origen")
    anchor = insert_after(install, "B.3. Reconstrucción de datos", "Heading 2")
    anchor = insert_after(
        anchor,
        "El dataset canónico no se publica en Git y debe reconstruirse. Primero se "
        "ejecutan 00_snapshot_legacy.py y 10_fetch_market_data.py. La consulta GDELT "
        "requiere un proyecto propio de Google Cloud con BigQuery y facturación; "
        "30_fetch_gdelt.py debe ejecutarse primero sin --execute para revisar el "
        "dry-run y el límite de bytes. La consulta sellada leyó 1.209,42 GiB. "
        "Después se ejecutan 35_audit_gdelt_candidates.py, 40_build_curated.py y "
        "50_compare_legacy_curated.py. La guía exacta está en docs/REPRODUCIR.md.",
    )
    insert_after(
        anchor,
        "El archivo resultante debe contener 2.558 filas y producir SHA-256 "
        "abf82d900314ce09cd00113785804e9f6742c38dd0867880ae9620c2a12f92ce antes "
        "de ejecutar modelos. No se debe lanzar la consulta GDELT sin revisar costo, "
        "cuota y presupuesto del proyecto utilizado.",
    )

    set_by_prefix(doc, "B.3. Estructura del pipeline", "B.4. Estructura del pipeline")
    set_by_prefix(doc, "B.4. Reproducción de los resultados", "B.5. Reproducción de los resultados")
    set_by_prefix(
        doc,
        "00_prepare.py  Preparación de datos",
        "00_prepare.py  Preparación de datos\n"
        "05_stationarity.py  ADF/KPSS y correlogramas\n"
        "10_baselines.py  Baseline RW\n"
        "11_train_arima.py  ARIMA iterado\n"
        "13_train_sarimax.py  ARIMAX/SARIMAX\n"
        "12_train_lstm.py  Variantes LSTM OOS\n"
        "14_walkforward.py  Walk-forward mensual\n"
        "15_multiseed_lstm.py  120 entrenamientos\n"
        "16_robustez_2025.py  Robustez 2025\n"
        "20_evaluate_stats.py / 24_wf_metrics.py  Evaluación\n"
        "30_make_figures.py / 31_regimes.py  Figuras y regímenes",
    )
    set_by_prefix(
        doc,
        "# Verificación rápida de mecánica",
        "# Pruebas\n"
        ".venv\\Scripts\\python.exe -m pytest -q\n\n"
        "# Verificación rápida aislada\n"
        ".venv\\Scripts\\python.exe run_all.py --quick\n\n"
        "# Reproducción completa con archivo de resultados previos\n"
        ".venv\\Scripts\\python.exe run_all.py --fresh",
    )
    set_by_prefix(
        doc,
        "El modo --quick crea una copia temporal",
        "El modo --quick crea una copia temporal aislada, reduce las LSTM a dos "
        "épocas y omite etapas costosas; no produce cifras de la tesis. --fresh "
        "archiva data/processed, outputs y reports en _run_archive/, recrea los "
        "directorios y ejecuta la cadena completa sin borrar resultados anteriores.",
    )
    set_by_prefix(
        doc,
        "Las salidas se escriben en outputs/preds/",
        "La corrida definitiva con --fresh terminó con RUN_ALL_OK en 38.0 minutos, "
        "el gate consolidado aprobó 17 controles y pytest aprobó 19 pruebas. Dos "
        "corridas independientes reprodujeron 512/512 predicciones y todas las tablas "
        "comparables byte a byte. El manifiesto final registra 17 tablas y 28 figuras, "
        "incluidos ADF/KPSS y los correlogramas reconstruidos.",
    )


def update_references(doc: Document) -> None:
    references = doc.tables[10]
    new_entries = [
        (
            "[30]",
            "Federal Reserve Bank of St. Louis, “NASDAQ-100 (NASDAQ100)”, FRED, "
            "fuente original Nasdaq, Inc. Accedido: 16 de julio de 2026. [En línea]. "
            "Disponible en: https://fred.stlouisfed.org/series/NASDAQ100",
        ),
        (
            "[31]",
            "Cboe Global Markets, “VIX Index Historical Data”. Accedido: 16 de julio "
            "de 2026. [En línea]. Disponible en: "
            "https://cdn.cboe.com/api/global/us_indices/daily_prices/VIX_History.csv",
        ),
        (
            "[32]",
            "The GDELT Project, “GDELT 2.0 Global Knowledge Graph”, conjunto público "
            "de Google BigQuery gdelt-bq.gdeltv2.gkg_partitioned. Accedido: 16 de "
            "julio de 2026. [En línea]. Disponible en: https://www.gdeltproject.org/",
        ),
    ]
    for number, reference in new_entries:
        row = references.add_row()
        set_cell_text(row.cells[0], number)
        set_cell_text(row.cells[1], reference)


def refine_pagination(doc: Document) -> None:
    table_3_caption = find_one(
        doc,
        "Tabla 3. Cuadro Comparativa de estudios relevantes",
        "Caption",
    )
    table_3_anchor = find_one(doc, "La presente tesis aborda ambos vacíos")
    remove_empty_between(table_3_anchor._p, table_3_caption._p)
    table_3_caption.paragraph_format.keep_with_next = True

    table_3_continuation = find_one(doc, "Continuación: Tabla 4.", "Caption")
    set_paragraph(
        table_3_continuation,
        "Continuación: Tabla 3. Cuadro Comparativa de estudios relevantes "
        "(orden cronológico)",
    )
    table_3_continuation.paragraph_format.keep_with_next = True
    for table in (doc.tables[3], doc.tables[4]):
        row_properties = table.rows[0]._tr.get_or_add_trPr()
        if row_properties.find(qn("w:tblHeader")) is None:
            row_properties.append(OxmlElement("w:tblHeader"))
        for row in table.rows:
            row_properties = row._tr.get_or_add_trPr()
            if row_properties.find(qn("w:cantSplit")) is None:
                row_properties.append(OxmlElement("w:cantSplit"))

    chapter_7 = find_one(doc, "RESULTADOS Y DISCUSIÓN", "Heading 1")
    chapter_6_end = find_one(doc, "Análisis por horizonte y lectura de aportes.")
    remove_empty_between(chapter_6_end._p, chapter_7._p)
    chapter_7.paragraph_format.page_break_before = True

    table_71_caption = find_one(doc, "Tabla 7.1 — Métricas fuera de muestra", "Caption")
    table_71_note = find_one(doc, "Mejor RMSE por horizonte sombreado en verde. ¹")
    remove_empty_between(table_71_caption._p, table_71_note._p)
    table_71_caption.paragraph_format.keep_with_next = True
    table_71_note.paragraph_format.keep_with_next = True

    discussion_heading = find_one(doc, "7.7 Discusión general", "Heading 2")
    discussion_caption = find_one(doc, "Figura 7.7", "Caption")
    discussion_picture = discussion_caption._p.getprevious()
    while discussion_picture is not None and not discussion_picture.xpath(".//a:blip"):
        discussion_picture = discussion_picture.getprevious()
    if discussion_picture is None:
        raise RuntimeError("No picture found before Figura 7.7")
    remove_empty_between(discussion_heading._p, discussion_picture)
    discussion_first = find_one(doc, "Tomados en conjunto, los resultados")
    remove_empty_between(discussion_caption._p, discussion_first._p)

    for prefix in (
        "7.4 Estabilidad temporal",
        "7.5 Sensibilidad a la inicialización",
        "REFERENCIAS",
    ):
        find_one(doc, prefix).paragraph_format.page_break_before = True
    chapter_8 = find_one(
        doc,
        "CONCLUSIONES, IMPLICACIONES Y TRABAJO FUTURO",
        "Heading 1",
    )
    chapter_8.paragraph_format.page_break_before = False
    discussion_last = find_one(doc, "La cuarta lectura es metodológica.")
    remove_empty_between(discussion_last._p, chapter_8._p)

    appendix = find_one(doc, "APÉNDICE", "Heading 1")
    remove_empty_between(doc.tables[10]._tbl, appendix._p)
    appendix.paragraph_format.page_break_before = False

    appendix_b = find_one(doc, "Apéndice B. Reproducibilidad", "Heading 1")
    pacf_caption = find_one(doc, "PACF – Precio NASDAQ-100 en nivel")
    remove_empty_between(pacf_caption._p, appendix_b._p)
    appendix_b.paragraph_format.page_break_before = True

    appendix_c = find_one(doc, "Apéndice C. Glosario técnico", "Caption")
    reproduction_end = find_one(doc, "La corrida definitiva con --fresh")
    remove_empty_between(reproduction_end._p, appendix_c._p)



def replace_figures(doc: Document) -> None:
    images = {
        "Figura 7.1": ROOT / "reports/figs/Fig_T20_dumbbell_dRMSE.png",
        "Figura 7.2": ROOT / "reports/figs/Fig_T20_volcano_vs_RW.png",
        "Figura 7.3": ROOT / "reports/figs/Fig_T20_calibracion_scatter.png",
        "Figura 7.4": ROOT / "reports/figs/Fig_T20_WF_heatmap_vs_RW.png",
        "Figura 7.5": ROOT / "reports/figs/Fig_T20_cumloss_vs_RW.png",
        "Figura 7.6": ROOT / "reports/figs/Fig_multiseed_boxplot.png",
        "Figura 7.7": ROOT / "reports/figs/Fig_bump_ranking.png",
        "ACF – Precio NASDAQ-100 en nivel": ROOT / "reports/figs/Fig_appendix_acf_level.png",
        "PACF – Precio NASDAQ-100 en nivel": ROOT / "reports/figs/Fig_appendix_pacf_level.png",
    }
    for caption, image in images.items():
        if not image.exists():
            raise FileNotFoundError(image)
        scale = 0.90 if caption == "Figura 7.7" else 1.0
        replace_picture_before_caption(doc, caption, image, scale)


def validate(doc: Document) -> None:
    forbidden = [
        "p = 0.302",
        "p = 0.048",
        "p = 0.038",
        "mediana de p-valores 0.190",
        "ensemble p = 0.136",
        "correcciones-auditoria-final",
        "Python 3.11",
        "RC2.1",
    ]
    body = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    for table in doc.tables:
        body += "\n" + "\n".join(cell.text for row in table.rows for cell in row.cells)
    leftovers = [term for term in forbidden if term in body]
    if leftovers:
        raise RuntimeError(f"Forbidden legacy claims remain: {leftovers}")
    if "abf82d900314ce09cd00113785804e9f6742c38dd0867880ae9620c2a12f92ce" not in body:
        raise RuntimeError("Canonical dataset hash is missing from the thesis")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    if args.source.resolve() == args.output.resolve():
        raise ValueError("Source and output must be different files")
    doc = Document(args.source)
    update_summary_and_abstract(doc)
    update_methodology(doc)
    update_tables(doc)
    update_results(doc)
    update_conclusions(doc)
    update_appendices(doc)
    update_references(doc)
    replace_figures(doc)
    refine_pagination(doc)
    validate(doc)
    doc.core_properties.title = "Tesis Maestro Final - Datos Reconstruidos"
    doc.core_properties.subject = "Resultados reproducidos sobre fuentes reconstruidas"
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
